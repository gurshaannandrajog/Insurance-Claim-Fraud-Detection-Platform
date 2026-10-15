"""Export a report dict to CSV, Excel (.xlsx) or PDF (REQ-REP-010). Same rows as the on-screen report."""
import csv
import io

from django.utils import timezone


def _flat_summary(summary):
    lines = []
    for k, v in summary.items():
        label = k.replace("_", " ").capitalize()
        if isinstance(v, list):
            parts = []
            for item in v:
                if "label" in item:
                    parts.append(f"{item['label']}: {item['value']}")
                elif "month" in item:
                    parts.append(f"{item['month']}: {item['value']}")
            lines.append((label, ", ".join(parts)))
        else:
            lines.append((label, f"{v:,.2f}" if isinstance(v, float) else str(v)))
    return lines


def _filters_line(filters):
    active = {k: v for k, v in filters.items() if v}
    return ", ".join(f"{k.replace('_', ' ')}={v}" for k, v in active.items()) or "none"


def to_csv(report, filters):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([report["title"]])
    w.writerow([f"Generated {timezone.localtime():%Y-%m-%d %H:%M}", f"Filters: {_filters_line(filters)}"])
    w.writerow([])
    w.writerow(report["columns"])
    w.writerows(report["rows"])
    w.writerow([])
    w.writerow(["Summary"])
    w.writerows(_flat_summary(report["summary"]))
    return ("﻿" + buf.getvalue()).encode("utf-8"), "text/csv; charset=utf-8", "csv"   # BOM → Excel opens UTF-8


def to_xlsx(report, filters):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.append([report["title"]])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Generated {timezone.localtime():%Y-%m-%d %H:%M}  ·  Filters: {_filters_line(filters)}"])
    ws.append([])
    ws.append(report["columns"])
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E3A5F")
        cell.alignment = Alignment(vertical="center")
    for row in report["rows"]:
        ws.append(row)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if report["rows"]:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(report['columns']))}{ws.max_row}"
    for i, col in enumerate(report["columns"], start=1):
        width = max([len(str(col))] + [len(str(r[i - 1])) for r in report["rows"][:300]])
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width + 2))
        if "(₹)" in col:
            for r in range(header_row + 1, ws.max_row + 1):
                ws.cell(row=r, column=i).number_format = "#,##0.00"
    s = wb.create_sheet("Summary")
    s.append(["Metric", "Value"])
    for c in s[1]:
        c.font = Font(bold=True)
    for line in _flat_summary(report["summary"]):
        s.append(list(line))
    s.column_dimensions["A"].width = 32
    s.column_dimensions["B"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"


def to_pdf(report, filters):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    def clean(v):  # the built-in PDF fonts have no ₹ glyph
        return str(v).replace("₹", "Rs.")

    styles = getSampleStyleSheet()
    small = styles["BodyText"].clone("small", fontSize=7, leading=8.5)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title=report["title"])
    story = [Paragraph(report["title"], styles["Title"]),
             Paragraph(f"Generated {timezone.localtime():%Y-%m-%d %H:%M} · Filters: {_filters_line(filters)} · "
                       f"{len(report['rows'])} row(s)", styles["Normal"]), Spacer(1, 4 * mm)]
    summary = [[Paragraph(f"<b>{clean(k)}</b>", small), Paragraph(clean(v), small)]
               for k, v in _flat_summary(report["summary"])]
    if summary:
        t = Table(summary, colWidths=[60 * mm, 200 * mm])
        t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey),
                               ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F1F5F9"))]))
        story += [t, Spacer(1, 5 * mm)]
    data = [[Paragraph(f"<b>{clean(c)}</b>", small) for c in report["columns"]]]
    data += [[Paragraph(clean(f"{v:,.2f}" if isinstance(v, float) else v), small) for v in row]
             for row in report["rows"]]
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A5F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(table if report["rows"] else Paragraph("No records match the selected filters.", styles["Normal"]))
    doc.build(story)
    return buf.getvalue(), "application/pdf", "pdf"


EXPORTERS = {"csv": to_csv, "xlsx": to_xlsx, "pdf": to_pdf}
